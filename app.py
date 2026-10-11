"""Consolidador de Notas Fiscais (PDF/DANFE + XML de NF-e) — versão em arquivo único.

Execute com:  streamlit run app.py
"""

from __future__ import annotations

import hashlib
import io
import pandas as pd
import pdfplumber
import re
import streamlit as st
import unicodedata
import xml.etree.ElementTree as ET
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import datetime
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ===========================================================================
# utils
# ===========================================================================
def remover_acentos(texto: str) -> str:
    """Remove acentos mantendo o comprimento do texto (útil para regex)."""
    nfkd = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in nfkd if not unicodedata.combining(c))


def normalizar(texto: str) -> str:
    """Texto em maiúsculas, sem acentos e com 'º/°' padronizados para 'O'."""
    texto = texto.replace("º", "o").replace("°", "o").replace("ª", "a")
    return remover_acentos(texto).upper()


def numero_nf_limpo(valor: str) -> str:
    """'000.012.345' -> '12345'. Mantém '0' se o número for zero."""
    digitos = re.sub(r"\D", "", valor or "")
    return digitos.lstrip("0") or ("0" if digitos else "")


def numero_nf_da_chave(chave: str) -> str | None:
    """Extrai o nNF (posições 26-34) de uma chave de acesso de 44 dígitos.

    Layout: cUF(2) AAMM(4) CNPJ(14) mod(2) série(3) nNF(9) tpEmis(1) cNF(8) cDV(1)
    """
    digitos = re.sub(r"\D", "", chave or "")
    if len(digitos) != 44:
        return None
    return numero_nf_limpo(digitos[25:34])


def chave_valida(chave: str) -> bool:
    """Valida o dígito verificador (módulo 11) de uma chave de acesso."""
    digitos = re.sub(r"\D", "", chave or "")
    if len(digitos) != 44:
        return False
    corpo, dv = digitos[:43], int(digitos[43])
    pesos = [2, 3, 4, 5, 6, 7, 8, 9]
    soma = sum(int(d) * pesos[i % 8] for i, d in enumerate(reversed(corpo)))
    resto = soma % 11
    esperado = 0 if resto < 2 else 11 - resto
    return dv == esperado


def converter_decimal(valor: str | None) -> float | None:
    """Converte números no formato brasileiro ou americano para float.

    '1.234,5600' -> 1234.56 | '10,5' -> 10.5 | '10.5000' -> 10.5 | '1.000' -> 1000
    """
    if valor is None:
        return None
    s = str(valor).strip().replace(" ", "")
    if not s:
        return None
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif s.count(".") > 1 or re.fullmatch(r"\d{1,3}\.\d{3}", s):
        # '1.234.567' ou '1.000' (separador de milhar sem decimais)
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


# ===========================================================================
# modelos
# ===========================================================================
SEM_REFERENCIA = "N/D"


@dataclass
class ItemNF:
    """Um item (linha de produto) de uma Nota Fiscal."""

    nf_referencia: str
    nf_emissao: str
    codigo_material: str
    quantidade: float | None
    descricao: str = ""
    unidade: str = ""
    data_emissao: str = ""  # ISO "AAAA-MM-DD" ("" se não encontrada)
    chave_acesso: str = ""
    arquivo: str = ""
    origem: str = ""  # "XML" ou "PDF"


@dataclass
class ResultadoArquivo:
    """Resultado do processamento de um arquivo enviado."""

    arquivo: str
    tipo: str
    itens: list[ItemNF] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)
    conferencias: list[str] = field(default_factory=list)
    erro: str | None = None

    @property
    def notas(self) -> set[str]:
        return {i.nf_emissao for i in self.itens}

    @property
    def status(self) -> str:
        if self.erro:
            return "Erro"
        if self.avisos:
            return "Com avisos"
        return "OK"


# ===========================================================================
# referencias
# ===========================================================================
# Palavra-chave que indica Nota Fiscal
_KW_NF = r"(?:NOTAS?\s+FISCA(?:L|IS)(?:\s+ELETRONICAS?)?|DANFE|\bN\.?\s?F\.?(?:\s?-?\s?E)?S?\b)"

# Qualificadores que podem aparecer entre a palavra-chave e o número
_QUALIF = (
    r"(?:\s*(?:DE|DA|DO|DAS|DOS|REF\.?|REFERENCIA(?:DA|DAS)?|REFERENTE|ORIGEM|ORIGINAL|"
    r"VENDA|REMESSA|RETORNO|DEVOLUCAO|ENTRADA|SAIDA|NUMEROS?|NUM\.?|NO\.?|NROS?\.?|N\.|N(?=\s*\d)|"
    r"SERIE\s*\d{1,3}\s*[,-]?|[:#\-.])\s*)*"
)

# Um número de NF (aceita pontos de milhar: 000.012.345) ...
_NUM = r"\d{1,3}(?:\.\d{3}){1,2}|\d{1,9}"
# ... ou uma lista: "10, 11 e 12" / "10/11" não (barra costuma ser série)
_LISTA = rf"(?:{_NUM})(?:\s*(?:,|;|\bE\b)\s*(?:{_NUM}))*"

_RE_NF = re.compile(rf"{_KW_NF}{_QUALIF}\s*({_LISTA})(?!\d)")
_RE_CHAVE = re.compile(r"(?<!\d)(?:\d[\s.]?){43}\d(?!\d)")
_RE_PEDIDO = re.compile(
    r"(?:\bPEDIDO(?:\s+DE\s+COMPRAS?)?|\bPED\.|ORDEM\s+DE\s+COMPRA|\bXPED)"
    r"\s*(?:N\.|NO\.?|NUM\.?|NUMERO|[:#\-])*\s*([A-Z0-9][A-Z0-9./-]{2,})"
)


def _sem_duplicatas(valores: Iterable[str]) -> list[str]:
    vistos: list[str] = []
    for v in valores:
        if v and v not in vistos:
            vistos.append(v)
    return vistos


def numeros_nf_no_texto(texto: str, ignorar: str | None = None) -> list[str]:
    """Retorna os números de NF citados no texto (sem zeros à esquerda)."""
    if not texto:
        return []
    t = normalizar(texto)
    encontrados: list[str] = []

    # 1) Chaves de acesso de 44 dígitos (mais confiáveis)
    for m in _RE_CHAVE.finditer(t):
        chave = re.sub(r"\D", "", m.group(0))
        if chave_valida(chave):
            encontrados.append(numero_nf_da_chave(chave))
    t_sem_chaves = _RE_CHAVE.sub(" ", t)

    # 2) Menções textuais: "NF 123", "Nota Fiscal nº 456"...
    for m in _RE_NF.finditer(t_sem_chaves):
        for parte in re.split(r"\s*(?:,|;|\bE\b)\s*", m.group(1)):
            numero = numero_nf_limpo(parte)
            if numero and numero != "0":
                encontrados.append(numero)

    ignorar_limpo = numero_nf_limpo(ignorar) if ignorar else None
    return [n for n in _sem_duplicatas(encontrados) if n != ignorar_limpo]


def pedido_no_texto(texto: str) -> str | None:
    if not texto:
        return None
    m = _RE_PEDIDO.search(normalizar(texto))
    return m.group(1).rstrip(".-/") if m else None


def extrair_referencia(
    texto_complementar: str,
    nf_emissao: str | None = None,
    refs_estruturadas: Iterable[str] = (),
    usar_pedido: bool = True,
) -> str:
    """Monta o valor da coluna "Nota Fiscal de Referência".

    Prioridade: referências estruturadas (tags NFref do XML) + NFs citadas no texto;
    se nenhuma for encontrada, usa o número do pedido (quando houver); senão "N/D".
    """
    proprio = numero_nf_limpo(nf_emissao) if nf_emissao else None
    numeros = [n for n in refs_estruturadas if n and n != proprio]
    numeros += numeros_nf_no_texto(texto_complementar, ignorar=nf_emissao)
    numeros = _sem_duplicatas(numeros)
    if numeros:
        return ", ".join(numeros)
    if usar_pedido:
        pedido = pedido_no_texto(texto_complementar)
        if pedido:
            return f"Pedido {pedido}"
    return SEM_REFERENCIA


# ===========================================================================
# xml_parser
# ===========================================================================
def _local(tag: str) -> str:
    """'{http://www.portalfiscal.inf.br/nfe}det' -> 'det'."""
    return tag.rsplit("}", 1)[-1]


def _filho(elem: ET.Element | None, *caminho: str) -> ET.Element | None:
    """Navega pelos filhos diretos ignorando namespaces."""
    atual = elem
    for nome in caminho:
        if atual is None:
            return None
        atual = next((c for c in atual if _local(c.tag) == nome), None)
    return atual


def _texto(elem: ET.Element | None, *caminho: str) -> str:
    alvo = _filho(elem, *caminho)
    return (alvo.text or "").strip() if alvo is not None else ""


def _referencias_estruturadas(ide: ET.Element | None) -> list[str]:
    """Números das NFs informadas nos grupos NFref (refNFe, refNF, refNFP...)."""
    numeros: list[str] = []
    if ide is None:
        return numeros
    for nfref in (c for c in ide if _local(c.tag) == "NFref"):
        for ref in nfref:
            nome = _local(ref.tag)
            if nome in ("refNFe", "refNFeSig", "refCTe"):
                n = numero_nf_da_chave(ref.text or "")
            elif nome in ("refNF", "refNFP"):
                n = numero_nf_limpo(_texto(ref, "nNF"))
            elif nome == "refECF":
                n = numero_nf_limpo(_texto(ref, "nCOO"))
            else:
                n = None
            if n:
                numeros.append(n)
    return numeros


def _quantidade(valor: str) -> float | None:
    # No XML o separador decimal é sempre ponto ("10.0000")
    try:
        return float(valor) if valor else None
    except ValueError:
        return None


def processar_xml(conteudo: bytes, nome_arquivo: str) -> ResultadoArquivo:
    resultado = ResultadoArquivo(arquivo=nome_arquivo, tipo="XML")
    try:
        raiz = ET.fromstring(conteudo)
    except ET.ParseError as exc:
        resultado.erro = f"XML inválido: {exc}"
        return resultado

    blocos = [e for e in raiz.iter() if _local(e.tag) == "infNFe"]
    if not blocos:
        nome_raiz = _local(raiz.tag)
        resultado.erro = (
            f"O arquivo não contém uma NF-e (elemento raiz: <{nome_raiz}>). "
            "XMLs de evento, cancelamento ou CT-e não são suportados."
        )
        return resultado

    for inf in blocos:
        ide = _filho(inf, "ide")
        nf_emissao = numero_nf_limpo(_texto(ide, "nNF"))
        if not nf_emissao:
            # fallback: atributo Id="NFe<chave>"
            nf_emissao = numero_nf_da_chave(inf.get("Id", "")) or ""
        if not nf_emissao:
            resultado.avisos.append("Número da NF (nNF) não encontrado.")

        # Data de emissão: dhEmi (NF-e 3.10/4.00, "2026-10-01T09:30:00-04:00") ou dEmi (2.00)
        data_emissao = (_texto(ide, "dhEmi") or _texto(ide, "dEmi"))[:10]
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", data_emissao):
            data_emissao = ""
            resultado.avisos.append(f"NF {nf_emissao or '?'}: data de emissão não encontrada no XML.")
        chave = re.sub(r"\D", "", inf.get("Id", ""))
        inf_cpl = _texto(inf, "infAdic", "infCpl")
        referencia = extrair_referencia(
            inf_cpl,
            nf_emissao=nf_emissao,
            refs_estruturadas=_referencias_estruturadas(ide),
        )

        dets = [c for c in inf if _local(c.tag) == "det"]
        if dets:
            resultado.conferencias.append(
                f"NF {nf_emissao or '?'}: {len(dets)} {'item' if len(dets) == 1 else 'itens'} (XML oficial ✔)"
            )
        if not dets:
            resultado.avisos.append(f"NF {nf_emissao or '?'}: nenhum item <det> encontrado.")
        for det in dets:
            prod = _filho(det, "prod")
            qtd_txt = _texto(prod, "qCom")
            quantidade = _quantidade(qtd_txt)
            if quantidade is None:
                resultado.avisos.append(
                    f"NF {nf_emissao}, item {det.get('nItem', '?')}: quantidade inválida ({qtd_txt!r})."
                )
            resultado.itens.append(
                ItemNF(
                    nf_referencia=referencia,
                    nf_emissao=nf_emissao,
                    codigo_material=_texto(prod, "cProd"),
                    quantidade=quantidade,
                    descricao=_texto(prod, "xProd"),
                    unidade=_texto(prod, "uCom"),
                    data_emissao=data_emissao,
                    chave_acesso=chave if len(chave) == 44 else "",
                    arquivo=nome_arquivo,
                    origem="XML",
                )
            )
    return resultado


# ===========================================================================
# pdf_parser
# ===========================================================================
_RE_CHAVE = re.compile(r"(?<!\d)(?:\d[ .]?){43}\d(?!\d)")

_RE_NUMERO_DANFE = [
    # "Nº 000.012.345" / "N. 000.012.345" (formato padrão do DANFE)
    re.compile(r"\bN\s*O?\s*[.:]?\s*(\d{3}\.\d{3}\.\d{3})\b"),
    re.compile(r"(?:NF-?E|NOTA\s+FISCAL(?:\s+ELETRONICA)?)\s*N\s*O?\s*[.:]?\s*(\d{1,9})\b"),
]

# Linha de item no layout padrão:
# CÓDIGO  DESCRIÇÃO  NCM/SH  O/CST  CFOP  UN  QUANT  V.UNIT  V.TOTAL ...
# A descrição é opcional: quando ela é longa, o DANFE a quebra em várias linhas
# e o código fica sozinho na linha dos números, com o texto acima/abaixo.
_NUM_BR = r"\d{1,3}(?:\.\d{3})*,\d+|\d+,\d+"
_RE_LINHA_ITEM = re.compile(
    r"^(?P<cod>\S+)(?:\s+(?P<desc>.+?))?\s+"
    r"(?P<ncm>\d{8}|\d{4}\.\d{2}\.\d{2})\s+"
    r"(?P<cst>\d{2,4})\s+"
    r"(?P<cfop>\d{4}|\d\.\d{3})\s+"
    r"(?P<un>[A-Za-z][\w/.-]{0,7})\s+"
    r"(?P<qtd>\d{1,3}(?:\.\d{3})+(?:,\d+)?|\d+(?:,\d+)?)\s+"
    rf"(?P<vunit>{_NUM_BR})\s+"
    rf"(?P<vtot>{_NUM_BR})"
)
# Linha que "parece" item (tem NCM de 8 dígitos e CFOP) mas não foi reconhecida
_RE_LINHA_SUSPEITA = re.compile(r"\b\d{8}\b.*\b[1-7]\d{3}\b")
_RE_TOTAL_PRODUTOS = re.compile(r"VALOR\s+TOTAL\s+DOS\s+PRODUTOS")
_RE_VALOR = re.compile(rf"(?<![\d,.])(?:{_NUM_BR})(?![\d,])")
_TOLERANCIA_TOTAL = 0.05  # R$

_RE_INFO_COMPL = re.compile(r"INFORMA[ÇC][ÕO]ES\s+COMPLEMENTARES|DADOS\s+ADICIONAIS", re.I)
_RE_RESERVADO = re.compile(r"RESERVADO\s+AO\s+FISCO", re.I)


@dataclass
class _Pagina:
    texto: str
    tabelas: list[list[list[str | None]]]
    info_compl: str
    chave: str | None
    numero: str  # número da NF lido do campo "Nº" (usado se não houver chave)
    itens_pos: list[dict] = field(default_factory=list)  # itens lidos pela posição das palavras
    data_emissao: str = ""  # ISO "AAAA-MM-DD"

    @property
    def ident(self) -> str:
        return self.chave or (f"NF:{self.numero}" if self.numero else "")


@dataclass
class _Nota:
    paginas: list[_Pagina] = field(default_factory=list)

    @property
    def chave(self) -> str | None:
        return next((p.chave for p in self.paginas if p.chave), None)

    @property
    def ident(self) -> str:
        return next((p.ident for p in self.paginas if p.ident), "")


# --------------------------------------------------------------------------- #
# Leitura das páginas
# --------------------------------------------------------------------------- #
def _chave_da_pagina(texto: str) -> str | None:
    for m in _RE_CHAVE.finditer(texto):
        chave = re.sub(r"\D", "", m.group(0))
        if chave[20:22] in ("55", "65") and chave_valida(chave):
            return chave
    return None


def _texto_info_complementar(page: pdfplumber.page.Page, texto_pagina: str) -> str:
    """Recorta o quadro "Informações Complementares" pela posição na página."""
    try:
        achados = page.search(_RE_INFO_COMPL.pattern, regex=True, case=False)
    except Exception:  # noqa: BLE001 - versões antigas do pdfplumber
        achados = []

    if achados:
        topo = min(a["top"] for a in achados)
        x_dir = page.width
        for r in page.search(_RE_RESERVADO.pattern, regex=True, case=False) or []:
            if r["top"] >= topo - 5 and r["x0"] > page.width * 0.3:
                x_dir = min(x_dir, r["x0"] - 2)
        recorte = page.crop((0, topo, x_dir, page.height), strict=False)
        texto = recorte.extract_text() or ""
        return _RE_INFO_COMPL.sub(" ", texto).strip()

    # Fallback textual: tudo após o título do quadro
    m = _RE_INFO_COMPL.search(texto_pagina)
    if not m:
        return ""
    trecho = texto_pagina[m.end():]
    return _RE_RESERVADO.sub(" ", trecho).strip()


def _sem_marca_dagua(page: pdfplumber.page.Page) -> pdfplumber.page.Page:
    """Remove caracteres de marca d'água: inclinados/rotacionados ou com fonte muito
    maior que a do texto normal (ex.: "NF-E SEM VALOR FISCAL" atravessando a página)."""
    tamanhos = sorted(c["size"] for c in page.chars)
    if not tamanhos:
        return page
    mediana = tamanhos[len(tamanhos) // 2]

    def manter(obj) -> bool:
        if obj.get("object_type") != "char":
            return True
        a, b, c, d, *_ = obj.get("matrix") or (1, 0, 0, 1, 0, 0)
        horizontal = abs(b) < 1e-3 and abs(c) < 1e-3 and a > 0 and d > 0
        return horizontal and obj["size"] <= mediana * 1.8

    return page.filter(manter)


# Parte numérica de uma linha de item (do NCM em diante) — usada por busca, não do início
_RE_CAUDA_ITEM = re.compile(
    r"(?:^|\s)(?P<ncm>\d{8}|\d{4}\.\d{2}\.\d{2})\s+"
    r"(?P<cst>\d{2,4})\s+"
    r"(?P<cfop>\d{4}|\d\.\d{3})\s+"
    r"(?P<un>[A-Za-z][\w/.-]{0,7})\s+"
    r"(?P<qtd>\d{1,3}(?:\.\d{3})+(?:,\d+)?|\d+(?:,\d+)?)\s+"
    rf"(?P<vunit>{_NUM_BR})\s+"
    rf"(?P<vtot>{_NUM_BR})"
)
# Títulos de quadro (bastam sozinhos) e títulos de coluna (precisam aparecer 2+ juntos)
_TITULOS_QUADRO = ("DADOS DO PRODUTO", "DADOS ADICIONAIS", "INFORMACOES COMPLEMENTARES",
                   "CALCULO DO ISSQN", "RESERVADO AO FISCO")
_TITULOS_COLUNA = ("CODIGO", "DESCRI", "NCM", "QUANT", "CFOP", "CST", "UNID", "VALOR UNIT", "V.UNIT")


def _linhas_de_palavras(page: pdfplumber.page.Page) -> list[dict]:
    """Agrupa as palavras em linhas pela posição vertical. Cada linha guarda as palavras
    (com coordenadas), o texto e o deslocamento de cada palavra dentro do texto."""
    palavras = page.extract_words(keep_blank_chars=False, use_text_flow=False, x_tolerance=1.5)
    palavras.sort(key=lambda w: (round(w["top"], 1), w["x0"]))
    linhas: list[dict] = []
    for w in palavras:
        if linhas and abs(w["top"] - linhas[-1]["top"]) <= 1.5:
            linhas[-1]["palavras"].append(w)
        else:
            linhas.append({"top": w["top"], "palavras": [w]})
    for l in linhas:
        l["palavras"].sort(key=lambda w: w["x0"])
        l["bottom"] = max(w["bottom"] for w in l["palavras"])
        l["meio"] = (l["top"] + l["bottom"]) / 2
        partes, offsets, pos = [], [], 0
        for w in l["palavras"]:
            offsets.append(pos)
            partes.append(w["text"])
            pos += len(w["text"]) + 1
        l["texto"] = " ".join(partes)
        l["offsets"] = offsets
    return linhas


def _eh_cabecalho(texto: str) -> bool:
    t = normalizar(texto)
    if any(k in t for k in _TITULOS_QUADRO):
        return True
    return sum(k in t for k in _TITULOS_COLUNA) >= 2


def _itens_por_posicao(page: pdfplumber.page.Page) -> list[dict]:
    """Lê os itens usando a posição das palavras (robusto a descrições/códigos que
    quebram em várias linhas, com alinhamento no topo, no meio ou na base da célula)."""
    linhas = _linhas_de_palavras(page)
    if not linhas:
        return []

    # 1) Âncoras: linhas que contêm a parte numérica do item
    ancoras = []
    for i, l in enumerate(linhas):
        m = _RE_CAUDA_ITEM.search(l["texto"])
        if not m:
            continue
        inicio = m.start("ncm")
        prefixo = [w for w, off in zip(l["palavras"], l["offsets"]) if off < inicio]
        ancoras.append({"idx": i, "m": m, "prefixo": prefixo})
    if not ancoras:
        return []

    # 2) Fronteira entre a coluna do código e a da descrição
    x_codigo = min(a["prefixo"][0]["x0"] for a in ancoras if a["prefixo"]) if any(a["prefixo"] for a in ancoras) else 0
    inicios_desc = sorted(a["prefixo"][1]["x0"] for a in ancoras if len(a["prefixo"]) >= 2)
    if inicios_desc:
        fronteira = inicios_desc[len(inicios_desc) // 2] - 1.0
    else:
        fronteira = max((a["prefixo"][0]["x1"] for a in ancoras if a["prefixo"]), default=x_codigo) + 1.0
    # início da coluna NCM (onde começam os números): nada de continuação passa daqui
    x_numeros = min(
        w["x0"] for a in ancoras
        for w, off in zip(linhas[a["idx"]]["palavras"], linhas[a["idx"]]["offsets"])
        if off == a["m"].start("ncm")
    )

    def continuacao_valida(k: int) -> bool:
        """Linha de continuação: tem texto na coluna da descrição e nada nas colunas numéricas."""
        ws = linhas[k]["palavras"]
        if any(w["x1"] > x_numeros + 1 for w in ws):
            return False
        return any(w["x0"] >= fronteira for w in ws)

    # 3) Distribui as linhas "soltas" entre as âncoras pelo maior espaço vertical
    altura = sorted(l["bottom"] - l["top"] for l in linhas)[len(linhas) // 2] or 6.0
    dono: dict[int, list[int]] = {a["idx"]: [] for a in ancoras}
    idx_ancoras = [a["idx"] for a in ancoras]

    def soltas(de: int, ate: int) -> list[int]:
        # entre dois itens da mesma página não existe cabeçalho: toda linha é conteúdo
        return list(range(de, ate))

    # antes da primeira âncora / depois da última: só linhas coladas (sem cabeçalhos)
    primeira, ultima = idx_ancoras[0], idx_ancoras[-1]
    k = primeira - 1
    while k >= 0 and not _eh_cabecalho(linhas[k]["texto"]) and continuacao_valida(k) and \
            linhas[k + 1]["top"] - linhas[k]["top"] <= altura * 1.8:
        dono[primeira].insert(0, k)
        k -= 1
    k = ultima + 1
    while k < len(linhas) and not _eh_cabecalho(linhas[k]["texto"]) and continuacao_valida(k) and \
            linhas[k]["top"] - linhas[k - 1]["top"] <= altura * 1.8:
        dono[ultima].append(k)
        k += 1
    # entre âncoras consecutivas: corta no maior intervalo vertical
    for a, b in zip(idx_ancoras, idx_ancoras[1:]):
        meio = [k for k in soltas(a + 1, b) if continuacao_valida(k) or
                all(w["x0"] < fronteira for w in linhas[k]["palavras"])]
        if not meio:
            continue
        seq = [a] + meio + [b]
        gaps = [linhas[seq[j + 1]]["top"] - linhas[seq[j]]["top"] for j in range(len(seq) - 1)]
        corte = max(range(len(gaps)), key=lambda j: gaps[j])  # depois de seq[corte]
        for j, k in enumerate(meio, start=1):
            (dono[a] if j <= corte else dono[b]).append(k)

    # 4) Monta cada item: código = palavras na coluna do código (de cima para baixo)
    itens = []
    for a in ancoras:
        m = a["m"]
        pedacos_cod: list[tuple[float, float, str]] = []
        pedacos_desc: list[tuple[float, float, str]] = []
        for w in a["prefixo"]:
            alvo = pedacos_cod if w["x0"] < fronteira else pedacos_desc
            alvo.append((w["top"], w["x0"], w["text"]))
        for k in dono[a["idx"]]:
            for w in linhas[k]["palavras"]:
                if w["x0"] < fronteira and w["x1"] >= x_codigo - 2:
                    pedacos_cod.append((w["top"], w["x0"], w["text"]))
                elif w["x0"] >= fronteira:
                    pedacos_desc.append((w["top"], w["x0"], w["text"]))
        # ordena por linha (top arredondado) e depois por x
        def juntar(pedacos, sep):
            pedacos.sort(key=lambda p: (round(p[0]), p[1]))
            return sep.join(p[2] for p in pedacos)

        codigo = juntar(pedacos_cod, "")
        if not codigo:
            continue
        itens.append(
            {
                "codigo": codigo,
                "descricao": juntar(pedacos_desc, " "),
                "unidade": m["un"],
                "quantidade": converter_decimal(m["qtd"]),
                "valor_total": converter_decimal(m["vtot"]),
            }
        )
    return itens


_RE_DATA = re.compile(r"(?<!\d)(\d{2})/(\d{2})/(\d{4})(?!\d)")
_RE_ROTULO_EMISSAO = re.compile(r"\b(?:DATA\s+D[AEO]\s+|DT\.?\s*(?:DE\s+)?)EMISSAO\b")


def _data_iso(dia: str, mes: str, ano: str) -> str:
    from datetime import date

    try:
        return date(int(ano), int(mes), int(dia)).isoformat()
    except ValueError:
        return ""


def _data_emissao_da_pagina(page: pdfplumber.page.Page, texto: str) -> str:
    """Data de emissão da NF no DANFE.

    1) Pela posição: a data logo ABAIXO do rótulo "DATA DE EMISSÃO" (quadro do
       destinatário), alinhada com ele — evita pegar a "DATA DE SAÍDA/ENTRADA".
    2) Na mesma linha, logo após o rótulo (ex.: canhoto "data de emissão: 01/10/2026").
    3) Pelo texto: primeira data nas 2 linhas seguintes ao rótulo.
    """
    linhas = _linhas_de_palavras(page)
    for i, l in enumerate(linhas):
        t = normalizar(l["texto"])
        m = _RE_ROTULO_EMISSAO.search(t)
        if not m:
            continue
        # palavras que formam o rótulo (pelo deslocamento no texto)
        rotulo = [w for w, off in zip(l["palavras"], l["offsets"]) if m.start() <= off < m.end()]
        if not rotulo:
            continue
        x0, x1 = rotulo[0]["x0"], rotulo[-1]["x1"]
        # 1) abaixo e alinhada
        for abaixo in linhas[i + 1 : i + 4]:
            if abaixo["top"] - l["bottom"] > 25:
                break
            for w in abaixo["palavras"]:
                d = _RE_DATA.search(w["text"])
                centro = (w["x0"] + w["x1"]) / 2
                if d and x0 - 20 <= centro <= x1 + 20:
                    iso = _data_iso(*d.groups())
                    if iso:
                        return iso
        # 2) na mesma linha, logo depois do rótulo
        depois = l["texto"][l["offsets"][l["palavras"].index(rotulo[-1])] :]
        d = _RE_DATA.search(depois[:40])
        if d and _data_iso(*d.groups()):
            return _data_iso(*d.groups())
    # 3) fallback textual
    linhas_txt = texto.splitlines()
    for i, linha in enumerate(linhas_txt):
        if _RE_ROTULO_EMISSAO.search(normalizar(linha)):
            for trecho in [linha] + linhas_txt[i + 1 : i + 3]:
                d = _RE_DATA.search(trecho)
                if d and _data_iso(*d.groups()):
                    return _data_iso(*d.groups())
    return ""


def _ler_paginas(pdf: pdfplumber.PDF) -> list[_Pagina]:
    paginas = []
    for page_original in pdf.pages:
        page = _sem_marca_dagua(page_original)
        texto = page.extract_text() or ""
        try:
            itens_pos = _itens_por_posicao(page)
        except Exception:  # noqa: BLE001 - nunca derrubar a leitura por causa desta estratégia
            itens_pos = []
        try:
            tabelas = page.extract_tables() or []
        except Exception:  # noqa: BLE001
            tabelas = []
        paginas.append(
            _Pagina(
                texto=texto,
                tabelas=tabelas,
                info_compl=_texto_info_complementar(page, texto) if texto else "",
                chave=_chave_da_pagina(texto),
                numero=_numero_no_texto(texto),
                itens_pos=itens_pos,
                data_emissao=_data_emissao_da_pagina(page, texto) if texto else "",
            )
        )
    return paginas


def _agrupar_por_nota(paginas: list[_Pagina]) -> list[_Nota]:
    """Páginas consecutivas da mesma nota (mesma chave ou, sem chave, mesmo Nº)
    são agrupadas; páginas sem identificação seguem a nota anterior."""
    notas: list[_Nota] = []
    for pag in paginas:
        atual = notas[-1].ident if notas else ""
        if not notas or (pag.ident and atual and pag.ident != atual and not _mesma_nota(pag, notas[-1])):
            notas.append(_Nota())
        notas[-1].paginas.append(pag)
    return notas


def _mesma_nota(pag: _Pagina, nota: _Nota) -> bool:
    """Compara pelo número quando só uma das partes tem a chave legível."""
    numero_nota = numero_nf_da_chave(nota.chave) if nota.chave else nota.ident.removeprefix("NF:")
    numero_pag = numero_nf_da_chave(pag.chave) if pag.chave else pag.numero
    if pag.chave and nota.chave:
        return pag.chave == nota.chave
    return bool(numero_pag) and numero_pag == numero_nota


# --------------------------------------------------------------------------- #
# Extração dos campos
# --------------------------------------------------------------------------- #
def _numero_no_texto(texto_pagina: str) -> str:
    texto = normalizar(texto_pagina)
    for padrao in _RE_NUMERO_DANFE:
        m = padrao.search(texto)
        if m:
            return numero_nf_limpo(m.group(1))
    return ""


def _numero_nf(nota: _Nota) -> str:
    if nota.chave:
        return numero_nf_da_chave(nota.chave) or ""
    return next((p.numero for p in nota.paginas if p.numero), "")


def _linha_de_texto_livre(linha: str) -> bool:
    """Linha que pode ser continuação de descrição (não é cabeçalho, marca d'água etc.)."""
    t = normalizar(linha)
    if len(t.replace(" ", "")) <= 2:  # letras soltas da marca d'água
        return False
    return not any(k in t for k in ("DESCRI", "DADOS ADICIONAIS", "DADOS DO PRODUTO", "INSCRI", "NATUREZA"))


def _itens_por_texto(nota: _Nota) -> tuple[list[dict], list[str]]:
    """Retorna (itens, linhas suspeitas que parecem item mas não foram lidas)."""
    itens: list[dict] = []
    suspeitas: list[str] = []
    for pag in nota.paginas:
        linhas = [l.strip() for l in pag.texto.splitlines()]
        casou = [_RE_LINHA_ITEM.match(l) for l in linhas]
        for i, (linha, m) in enumerate(zip(linhas, casou)):
            if not m:
                if _RE_LINHA_SUSPEITA.search(linha) and not _RE_CHAVE.search(linha):
                    suspeitas.append(linha)
                continue
            descricao = (m["desc"] or "").strip()
            if not descricao:
                # descrição quebrada: junta a linha anterior e a seguinte (se forem texto livre)
                partes = []
                if i > 0 and not casou[i - 1] and _linha_de_texto_livre(linhas[i - 1]):
                    partes.append(linhas[i - 1])
                if i + 1 < len(linhas) and not casou[i + 1] and _linha_de_texto_livre(linhas[i + 1]) \
                        and not _RE_LINHA_SUSPEITA.search(linhas[i + 1]):
                    partes.append(linhas[i + 1])
                descricao = " ".join(partes)
            itens.append(
                {
                    "codigo": m["cod"],
                    "descricao": descricao,
                    "unidade": m["un"],
                    "quantidade": converter_decimal(m["qtd"]),
                    "valor_total": converter_decimal(m["vtot"]),
                }
            )
    return itens, suspeitas


def _total_produtos(nota: _Nota) -> float | None:
    """Lê o campo "VALOR TOTAL DOS PRODUTOS" do quadro Cálculo do Imposto."""
    for pag in nota.paginas:
        linhas = pag.texto.splitlines()
        for i, linha in enumerate(linhas):
            if _RE_TOTAL_PRODUTOS.search(normalizar(linha)):
                # o título da coluna é o último da linha; o valor é o último número da linha seguinte
                for prox in linhas[i + 1 : i + 3]:
                    valores = _RE_VALOR.findall(prox)
                    if valores:
                        return converter_decimal(valores[-1])
    return None


def _soma(itens: list[dict]) -> float | None:
    if not itens or any(i.get("valor_total") is None for i in itens):
        return None
    return sum(i["valor_total"] for i in itens)


def _escolher_estrategia(candidatos: list[list[dict]], total: float | None) -> list[dict]:
    """Escolhe o resultado mais confiável entre as estratégias (em ordem de preferência):
    1º) a que "fecha" com o total impresso da nota; 2º) a que leu mais itens."""
    if total is not None:
        for c in candidatos:
            s = _soma(c)
            if s is not None and abs(s - total) <= _TOLERANCIA_TOTAL:
                return c
    return max(candidatos, key=len)  # max() mantém o primeiro em caso de empate


def _codigos_fora_do_padrao(codigos: list[str]) -> list[str]:
    """Códigos muito mais longos que o normal da nota (sinal de texto "grudado" no código)."""
    if len(codigos) < 3:
        return [c for c in codigos if len(c) > 30]
    tamanhos = sorted(len(c) for c in codigos)
    mediana = tamanhos[len(tamanhos) // 2]
    return list(dict.fromkeys(c for c in codigos if len(c) > max(mediana * 1.8, mediana + 6)))


def _qtd_itens(n: int) -> str:
    return f"{n} item" if n == 1 else f"{n} itens"


def _brl(valor: float) -> str:
    return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _indice_coluna(cabecalho: list[str], *termos: str, excluir: tuple[str, ...] = ()) -> int | None:
    for i, cel in enumerate(cabecalho):
        if any(t in cel for t in termos) and not any(e in cel for e in excluir):
            return i
    return None


def _itens_por_tabela(nota: _Nota) -> list[dict]:
    itens = []
    for pag in nota.paginas:
        for tabela in pag.tabelas:
            col_cod = col_qtd = col_desc = col_un = None
            for linha in tabela:
                celulas = [(c or "").strip() for c in linha]
                normal = [normalizar(c).replace("\n", " ") for c in celulas]

                # Procura (ou reencontra) o cabeçalho do quadro de produtos
                cod = _indice_coluna(normal, "COD", excluir=("CODIGO NCM",))
                qtd = _indice_coluna(normal, "QUANT", "QTD", "QTDE")
                if cod is not None and qtd is not None and cod != qtd:
                    col_cod, col_qtd = cod, qtd
                    col_desc = _indice_coluna(normal, "DESCRI")
                    col_un = next((i for i, c in enumerate(normal) if c in ("UN", "UNID", "UND", "UNID.")), None)
                    continue
                if col_cod is None or max(col_cod, col_qtd) >= len(celulas):
                    continue

                codigos = [c for c in celulas[col_cod].split("\n") if c.strip()]
                qtds = [q for q in celulas[col_qtd].split("\n") if q.strip()]
                if not codigos or len(codigos) != len(qtds):
                    continue
                descs = celulas[col_desc].split("\n") if col_desc is not None else []
                uns = celulas[col_un].split("\n") if col_un is not None and col_un < len(celulas) else []
                for i, (c, q) in enumerate(zip(codigos, qtds)):
                    quantidade = converter_decimal(q)
                    if quantidade is None:
                        continue
                    itens.append(
                        {
                            "codigo": c.strip(),
                            # descrições podem quebrar linha; só usamos se o número bater
                            "descricao": descs[i].strip() if len(descs) == len(codigos) else "",
                            "unidade": uns[i].strip() if len(uns) == len(codigos) else "",
                            "quantidade": quantidade,
                        }
                    )
    return itens


# --------------------------------------------------------------------------- #
# Ponto de entrada
# --------------------------------------------------------------------------- #
def processar_pdf(conteudo: bytes, nome_arquivo: str) -> ResultadoArquivo:
    resultado = ResultadoArquivo(arquivo=nome_arquivo, tipo="PDF")
    try:
        with pdfplumber.open(io.BytesIO(conteudo)) as pdf:
            paginas = _ler_paginas(pdf)
    except Exception as exc:  # noqa: BLE001
        resultado.erro = f"Não foi possível abrir o PDF: {exc}"
        return resultado

    if not any(p.texto.strip() for p in paginas):
        resultado.erro = (
            "O PDF não possui texto selecionável (provavelmente é uma imagem escaneada). "
            "Envie o XML da nota ou um PDF gerado digitalmente."
        )
        return resultado

    for nota in _agrupar_por_nota(paginas):
        nf_emissao = _numero_nf(nota)
        if not nf_emissao:
            resultado.avisos.append("Número da NF não identificado em uma das notas do PDF.")

        info_compl = "\n".join(p.info_compl for p in nota.paginas if p.info_compl)
        referencia = extrair_referencia(info_compl, nf_emissao=nf_emissao)
        data_emissao = next((p.data_emissao for p in nota.paginas if p.data_emissao), "")
        if not data_emissao:
            resultado.avisos.append(f"NF {nf_emissao or '?'}: data de emissão não encontrada no PDF.")

        por_posicao = [it for p in nota.paginas for it in p.itens_pos]
        por_texto, suspeitas = _itens_por_texto(nota)
        por_tabela = _itens_por_tabela(nota)
        total = _total_produtos(nota)
        itens = _escolher_estrategia([por_posicao, por_texto, por_tabela], total)

        # ---- Conferência interna de completude ---------------------------------
        # O valor em R$ NÃO vai para o relatório: é usado só como "prova real" de que
        # nenhuma linha de item foi pulada (soma das linhas lidas = total impresso).
        rotulo = f"NF {nf_emissao or '?'}"
        soma = _soma(itens)
        if total is not None and soma is not None and itens:
            if abs(soma - total) <= _TOLERANCIA_TOTAL:
                resultado.conferencias.append(
                    f"{rotulo}: {_qtd_itens(len(itens))} — todos os itens lidos ✔"
                )
            else:
                resultado.avisos.append(
                    f"{rotulo}: ATENÇÃO — pode haver item faltando nesta nota "
                    f"({_qtd_itens(len(itens))} lidos). Confira com o PDF ou use o XML."
                )
        elif itens:
            resultado.conferencias.append(f"{rotulo}: {_qtd_itens(len(itens))} lidos (conferência automática indisponível neste layout)")
        lidos_ok = total is not None and soma is not None and abs(soma - total) <= _TOLERANCIA_TOTAL
        for cod in _codigos_fora_do_padrao([i["codigo"] for i in itens]):
            resultado.avisos.append(f"{rotulo}: código com formato incomum, confira: {cod!r}")
        for linha in ([] if lidos_ok else suspeitas[:5]):
            resultado.avisos.append(f"{rotulo}: linha não reconhecida como item: {linha[:90]!r}")

        if not itens:
            resultado.avisos.append(
                f"NF {nf_emissao or '?'}: nenhum item reconhecido no PDF. "
                "O layout pode ser diferente do padrão — prefira o XML."
            )
        for it in itens:
            if it["quantidade"] is None:
                resultado.avisos.append(f"NF {nf_emissao}: quantidade ilegível no item {it['codigo']}.")
            resultado.itens.append(
                ItemNF(
                    nf_referencia=referencia,
                    nf_emissao=nf_emissao,
                    codigo_material=it["codigo"],
                    quantidade=it["quantidade"],
                    descricao=it["descricao"],
                    unidade=it["unidade"],
                    data_emissao=data_emissao,
                    chave_acesso=nota.chave or "",
                    arquivo=nome_arquivo,
                    origem="PDF",
                )
            )
    return resultado


# ===========================================================================
# processador
# ===========================================================================
COLUNAS = {
    "data_emissao": "Data de Emissão da Nota",
    "nf_referencia": "Nota de Referência",
    "nf_emissao": "Nota Emitida",
    "codigo_material": "Código do Produto",
    "quantidade": "Quantidade",
    "descricao": "Descrição",
    "unidade": "Unidade",
    "chave_acesso": "Chave de Acesso",
    "arquivo": "Arquivo",
    "origem": "Origem",
}
# Colunas do relatório, na ordem pedida
COLUNAS_PRINCIPAIS = [
    "Data de Emissão da Nota",
    "Código do Produto",
    "Quantidade",
    "Nota Emitida",
    "Nota de Referência",
]


def detectar_tipo(nome: str, conteudo: bytes) -> str | None:
    """Identifica o tipo pelo conteúdo (assinatura) e, em último caso, pela extensão."""
    inicio = conteudo[:1024].lstrip(b"\xef\xbb\xbf \t\r\n")
    if inicio.startswith(b"%PDF"):
        return "pdf"
    if inicio.startswith(b"<"):
        return "xml"
    ext = nome.lower().rsplit(".", 1)[-1]
    return ext if ext in ("pdf", "xml") else None


def processar_arquivo(nome: str, conteudo: bytes) -> ResultadoArquivo:
    tipo = detectar_tipo(nome, conteudo)
    try:
        if tipo == "xml":
            return processar_xml(conteudo, nome)
        if tipo == "pdf":
            return processar_pdf(conteudo, nome)
    except Exception as exc:  # noqa: BLE001 - nunca derrubar o lote inteiro
        return ResultadoArquivo(arquivo=nome, tipo=(tipo or "?").upper(), erro=f"Erro inesperado: {exc}")
    return ResultadoArquivo(arquivo=nome, tipo="?", erro="Formato não suportado (envie PDF ou XML).")


def processar_arquivos(arquivos: Iterable[tuple[str, bytes]]) -> list[ResultadoArquivo]:
    return [processar_arquivo(nome, conteudo) for nome, conteudo in arquivos]


def consolidar(resultados: Iterable[ResultadoArquivo]) -> pd.DataFrame:
    """Une os itens de todos os arquivos em um único DataFrame."""
    linhas = [vars(item) for r in resultados for item in r.itens]
    df = pd.DataFrame(linhas, columns=list(COLUNAS.keys())).rename(columns=COLUNAS)
    df["Quantidade"] = pd.to_numeric(df["Quantidade"], errors="coerce")
    return df


def remover_duplicados(df: pd.DataFrame) -> tuple[pd.DataFrame, int]:
    """Remove notas enviadas em duplicidade (ex.: PDF e XML da mesma NF).

    Quando a mesma NF aparece em mais de um arquivo, mantém apenas o primeiro
    arquivo (dando preferência ao XML, que é a fonte oficial).
    """
    if df.empty:
        return df, 0
    # Identificador da nota: chave de acesso (única no país); sem chave, o número da NF
    ident = df["Chave de Acesso"].where(df["Chave de Acesso"] != "", "NF:" + df["Nota Emitida"])
    ordenado = df.assign(_id=ident, _prio=(df["Origem"] != "XML").astype(int)).sort_values("_prio", kind="stable")
    arquivo_escolhido = ordenado.groupby("_id", sort=False)["Arquivo"].first()
    manter = df["Arquivo"] == ident.map(arquivo_escolhido)
    manter |= df["Nota Emitida"].eq("") & df["Chave de Acesso"].eq("")
    return df[manter].reset_index(drop=True), int((~manter).sum())


def agrupar_quantidades(df: pd.DataFrame) -> pd.DataFrame:
    """Consolida: une as linhas com o mesmo código dentro da mesma nota e soma as quantidades.

    Chave = Data + Nota de Referência + Nota Emitida + Código do Produto. A ordem segue a
    primeira aparição de cada código na nota. Funciona com o DataFrame completo ou só
    com as colunas principais.
    """
    if df.empty:
        return df
    chaves = [c for c in ("Data de Emissão da Nota", "Nota de Referência", "Nota Emitida", "Código do Produto")
              if c in df.columns]
    regras = {"Quantidade": "sum"}
    for col in df.columns:
        if col in chaves or col == "Quantidade":
            continue
        if col in ("Arquivo", "Origem"):
            regras[col] = lambda s: ", ".join(dict.fromkeys(s))
        else:
            regras[col] = "first"
    out = df.groupby(chaves, as_index=False, sort=False, dropna=False).agg(regras)[list(df.columns)]
    out["Quantidade"] = out["Quantidade"].round(6)  # evita resíduos de ponto flutuante (0,1+0,2)
    return out


# ===========================================================================
# exportacao
# ===========================================================================
def para_excel(df: pd.DataFrame, abas_extras: dict[str, pd.DataFrame] | None = None,
               nome_aba: str = "Consolidado") -> bytes:
    """Gera um .xlsx: a primeira aba recebe `df`; `abas_extras` vira abas adicionais."""
    buffer = io.BytesIO()
    abas = {nome_aba: df, **(abas_extras or {})}
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for nome, tabela in abas.items():
            preparar_exportacao(tabela).to_excel(writer, sheet_name=nome[:31], index=False)

        for ws in writer.book.worksheets:
            for cel in ws[1]:
                cel.font = Font(bold=True, color="FFFFFF")
                cel.fill = PatternFill("solid", fgColor="1F4E78")
                cel.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, coluna in enumerate(ws.iter_cols(values_only=True), start=1):
                largura = max(len(str(v)) if v is not None else 0 for v in coluna)
                ws.column_dimensions[get_column_letter(i)].width = min(max(largura + 2, 10), 60)

        # Todas as células no formato "Geral" do Excel (nada de formato Texto)...
        for ws in writer.book.worksheets:
            for linha in ws.iter_rows():
                for cel in linha:
                    cel.number_format = "General"
            # ...exceto a data: é gravada como DATA de verdade (filtra/ordena como data).
            # No formato Geral o Excel mostraria o número serial (ex.: 46296).
            cabecalho = [c.value for c in ws[1]]
            if COL_DATA in cabecalho:
                col = cabecalho.index(COL_DATA) + 1
                for (cel,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                    if cel.value is not None:
                        cel.number_format = "DD/MM/YYYY"
                ws.column_dimensions[get_column_letter(col)].width = max(
                    ws.column_dimensions[get_column_letter(col)].width or 0, 14)
    return buffer.getvalue()


COL_DATA = "Data de Emissão da Nota"


def _data_como_data(valor):
    """'2026-10-01' -> date(2026, 10, 1); vazio -> None."""
    from datetime import date

    if isinstance(valor, str) and len(valor) == 10:
        try:
            return date.fromisoformat(valor)
        except ValueError:
            return None
    return None


def data_br(valor) -> str:
    """'2026-10-01' -> '01/10/2026' (para tela e CSV)."""
    d = _data_como_data(valor)
    return d.strftime("%d/%m/%Y") if d else ""


def _nf_como_numero(valor):
    """'6265' -> 6265 (evita o triângulo verde de "número armazenado como texto")."""
    return int(valor) if isinstance(valor, str) and valor.isdigit() else valor


def _codigo_como_numero(valor):
    """Código só com dígitos vira número ('12345' -> 12345).

    Códigos com zero à esquerda ('000123') continuam como texto, senão o Excel
    apagaria os zeros e o código ficaria errado. Códigos com letras/traço ('0402-81')
    são texto por natureza e já ficam corretos no formato Geral.
    """
    if isinstance(valor, str) and valor.isdigit() and not (len(valor) > 1 and valor.startswith("0")):
        return int(valor)
    return valor


def preparar_exportacao(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    if COL_DATA in out.columns:
        out[COL_DATA] = pd.Series([_data_como_data(v) for v in out[COL_DATA]], index=out.index, dtype=object)
    for col in ("Nota de Referência", "Nota Emitida"):
        if col in out.columns:
            out[col] = out[col].map(_nf_como_numero).astype(object)
    if "Código do Produto" in out.columns:
        out["Código do Produto"] = out["Código do Produto"].map(_codigo_como_numero).astype(object)
    if "Quantidade" in out.columns:
        # lista + dtype object: impede o pandas de converter tudo de volta para float
        out["Quantidade"] = pd.Series(
            [int(q) if isinstance(q, float) and q.is_integer() else q for q in out["Quantidade"]],
            index=out.index,
            dtype=object,
        )
    return out


def para_csv(df: pd.DataFrame) -> bytes:
    """CSV no padrão do Excel em português: separador ';', decimal ',' e BOM UTF-8."""
    out = preparar_exportacao(df)
    if COL_DATA in out.columns:
        out[COL_DATA] = out[COL_DATA].map(lambda d: d.strftime("%d/%m/%Y") if d else "")
    if "Quantidade" in out.columns:
        out["Quantidade"] = out["Quantidade"].map(
            lambda q: "" if q is None or q != q else str(q).replace(".", ",")
        )
    return out.to_csv(index=False, sep=";").encode("utf-8-sig")


# ===========================================================================
# Interface Streamlit
# ===========================================================================
st.set_page_config(page_title="Consolidador de NF-e", page_icon="🧾", layout="wide")


# --------------------------------------------------------------------------- #
# Processamento (com cache: reenviar o mesmo arquivo não reprocessa)
# --------------------------------------------------------------------------- #
@st.cache_data(show_spinner=False, max_entries=2000)
def _processar_em_cache(nome: str, digest: str, _conteudo: bytes) -> ResultadoArquivo:
    # O cache é indexado pelo nome + hash SHA-256 do arquivo.
    return processar_arquivo(nome, _conteudo)


def processar_uploads(arquivos) -> list[ResultadoArquivo]:
    resultados = []
    barra = st.progress(0.0, text="Processando arquivos...")
    for i, arq in enumerate(arquivos, start=1):
        conteudo = arq.getvalue()
        digest = hashlib.sha256(conteudo).hexdigest()
        resultados.append(_processar_em_cache(arq.name, digest, conteudo))
        barra.progress(i / len(arquivos), text=f"Processando {i}/{len(arquivos)}: {arq.name}")
    barra.empty()
    return resultados


def tabela_situacao(resultados: list[ResultadoArquivo]) -> pd.DataFrame:
    linhas = []
    for r in resultados:
        if r.erro:
            situacao = f"❌ {r.erro}"
        elif r.avisos:
            situacao = "⚠️ " + " | ".join(r.avisos)
        else:
            situacao = " | ".join(r.conferencias) or "✔"
        linhas.append({"Arquivo": r.arquivo, "Itens": len(r.itens), "Situação": situacao})
    return pd.DataFrame(linhas)


# --------------------------------------------------------------------------- #
# Interface
# --------------------------------------------------------------------------- #
st.title("🧾 Consolidador de Notas Fiscais")
st.caption(
    "Envie DANFEs (PDF) e/ou XMLs de NF-e. O resultado traz, item a item: "
    "Data de Emissão, Código do Produto, Quantidade, Nota Emitida e Nota de Referência."
)

arquivos = st.file_uploader(
    "Selecione ou arraste os arquivos",
    type=["pdf", "xml"],
    accept_multiple_files=True,
    help="Você pode selecionar vários arquivos de uma vez.",
)

if not arquivos:
    st.info("Aguardando arquivos PDF ou XML.")
    st.stop()

resultados = processar_uploads(arquivos)
situacao = tabela_situacao(resultados)
# Se o PDF e o XML da mesma nota forem enviados juntos, os itens contam uma vez só (XML tem prioridade).
itens, _ = remover_duplicados(consolidar(resultados))
itens = itens[COLUNAS_PRINCIPAIS]
# Consolidação: mesmo código na mesma nota vira uma linha só, com as quantidades somadas
df = agrupar_quantidades(itens)

# ---- Alertas -----------------------------------------------------------------
com_problema = [r for r in resultados if r.erro or r.avisos]
if com_problema:
    st.warning(
        f"⚠️ {len(com_problema)} arquivo(s) precisam de conferência: "
        + ", ".join(r.arquivo for r in com_problema)
        + ". Veja os detalhes em “Situação por arquivo”."
    )
else:
    st.success(f"✔ {len(resultados)} arquivo(s) lidos — {len(itens)} linhas de itens, {len(df)} códigos consolidados.")

with st.expander("Situação por arquivo", expanded=bool(com_problema)):
    st.dataframe(
        situacao,
        hide_index=True,
        width="stretch",
        column_config={"Situação": st.column_config.TextColumn(width="large")},
    )

# ---- Tabela consolidada ----------------------------------------------------
st.subheader("Consolidado por código")

if df.empty:
    st.error("Nenhum item foi extraído dos arquivos enviados.")
    st.stop()


def filtrar(tabela: pd.DataFrame, termo: str) -> pd.DataFrame:
    if not termo:
        return tabela
    termo = termo.strip().lower()
    texto = tabela.astype(str)
    if "Data de Emissão da Nota" in tabela.columns:  # permite buscar por "01/10/2026"
        texto["Data de Emissão da Nota"] = tabela["Data de Emissão da Nota"].map(data_br)
    mascara = texto.apply(lambda col: col.str.lower().str.contains(termo, regex=False)).any(axis=1)
    return tabela[mascara]


busca = st.text_input("Filtrar", placeholder="Digite um número de NF ou código de material...")
visao = filtrar(df, busca)
visao_itens = filtrar(itens, busca)


def mostrar(tabela: pd.DataFrame) -> None:
    # exibe 2340 em vez de 2340.0 (fracionários como 2,5 continuam com casas decimais)
    exibir = preparar_exportacao(tabela)
    if "Data de Emissão da Nota" in exibir.columns:
        exibir["Data de Emissão da Nota"] = tabela["Data de Emissão da Nota"].map(data_br)
    if "Quantidade" in exibir.columns:  # 2,5 com vírgula, como no Excel
        exibir["Quantidade"] = exibir["Quantidade"].map(lambda q: "" if q is None or q != q else str(q).replace(".", ","))
    st.dataframe(exibir.astype(str), hide_index=True, width="stretch")


mostrar(visao)
unificadas = len(itens) - len(df)
st.caption(
    f"{len(visao)} de {len(df)} linha(s). "
    + (f"{unificadas} linha(s) com código repetido foram unificadas e somadas." if unificadas else
       "Nenhum código repetido dentro das notas.")
)

with st.expander("Ver item a item (exatamente como está nas notas)"):
    mostrar(visao_itens)
    st.caption(f"{len(visao_itens)} de {len(itens)} linha(s).")

# ---- Downloads -------------------------------------------------------------
carimbo = datetime.now().strftime("%Y%m%d_%H%M")
d1, d2, _ = st.columns([1, 1, 3])
d1.download_button(
    "⬇️ Baixar Excel (.xlsx)",
    data=para_excel(visao, {"Item a item": visao_itens}),
    file_name=f"notas_consolidadas_{carimbo}.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    type="primary",
    on_click="ignore",
    width="stretch",
    help="Aba 'Consolidado' (códigos somados) + aba 'Item a item' (como na nota).",
)
d2.download_button(
    "⬇️ Baixar CSV (consolidado)",
    data=para_csv(visao),
    file_name=f"notas_consolidadas_{carimbo}.csv",
    mime="text/csv",
    on_click="ignore",
    width="stretch",
)
